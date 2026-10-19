"""Logbook entry editor.

Build a per-day `entries[]` list for the logbook auto-fill console script,
then export a ready-to-paste script.js (also copied to clipboard).

Run:
    python logbook_ui.py
"""

from __future__ import annotations

import datetime
import json
import subprocess
import sys
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from tkinter import font as tkfont

import openpyxl

DEFAULT_IN_TIME = "09:00"
DEFAULT_OUT_TIME = "05:00"
DEFAULT_IN_AMPM = "am"
DEFAULT_OUT_AMPM = "pm"
AMPM = ["am", "pm"]

WEEKDAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

# Company timesheet layout: data rows start at row 14, with NO/DAY/DATE/TASK/
# PROJECT NAME/DETAIL/S-O NO/IN/OUT/TOTAL in columns A-J.
TIMESHEET_DATA_START_ROW = 14
COL_NO = 1
COL_DAY = 2
COL_TASK = 4
COL_DETAIL = 6
COL_IN = 8
COL_OUT = 9

LIGHT = {
    "bg": "#f2f3f5",
    "surface": "#ffffff",
    "surface_alt": "#f7f8fa",
    "input": "#ffffff",
    "border": "#d9dce3",
    "text": "#14161a",
    "dim": "#6b7280",
    "faint": "#c4c8d0",
    "accent": "#4f46e5",
    "accent_hover": "#4338ca",
    "accent_fg": "#ffffff",
    "chip": "#eceef3",
    "off_bg": "#eceef2",
    "skip_bg": "#fff6e5",
    "skip_fg": "#b45309",
    "sun_bg": "#fdeeee",
    "sun_fg": "#b91c1c",
}

DARK = {
    "bg": "#17181c",
    "surface": "#1f2126",
    "surface_alt": "#23252b",
    "input": "#2a2c33",
    "border": "#363941",
    "text": "#eef0f2",
    "dim": "#9aa1ab",
    "faint": "#4a4e57",
    "accent": "#7c7ff2",
    "accent_hover": "#8c8ff5",
    "accent_fg": "#101114",
    "chip": "#30333b",
    "off_bg": "#1a1b1f",
    "skip_bg": "#2c2517",
    "skip_fg": "#f0b35a",
    "sun_bg": "#2e1f22",
    "sun_fg": "#f08a8a",
}

# (header, min width px, weight). The header strip and every row frame use the
# same column config, so cells line up without sharing one grid.
COLUMNS = [
    ("Day", 96, 0),
    ("Off", 64, 0),
    ("Skip", 70, 0),
    ("Clock In", 140, 0),
    ("Clock Out", 140, 0),
    ("Activity", 0, 1),
    ("Description", 0, 1),
]
CELL_PADX = 6


def configure_columns(frame: tk.Widget) -> None:
    for col, (_, minsize, weight) in enumerate(COLUMNS):
        if weight:
            frame.grid_columnconfigure(col, weight=weight, uniform="text")
        else:
            frame.grid_columnconfigure(col, minsize=minsize)


def detect_dark_mode() -> bool:
    """Follow the OS light/dark setting (macOS and Windows; light elsewhere)."""
    if sys.platform == "darwin":
        try:
            out = subprocess.run(
                ["defaults", "read", "-g", "AppleInterfaceStyle"],
                capture_output=True, text=True, timeout=2,
            )
            return out.stdout.strip().lower() == "dark"
        except (OSError, subprocess.SubprocessError):
            return False
    if sys.platform == "win32":
        try:
            import winreg

            key = winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                r"Software\Microsoft\Windows\CurrentVersion\Themes\Personalize",
            )
            value, _ = winreg.QueryValueEx(key, "AppsUseLightTheme")
            return value == 0
        except OSError:
            return False
    return False


def _is_zero_time(value: object) -> bool:
    """True for empty cells and 0:00-style values (non-working days)."""
    if value is None:
        return True
    if isinstance(value, datetime.time):
        return value.hour == 0 and value.minute == 0
    text = str(value).strip()
    return text in ("", "0", "0:00", "00:00", "0:00:00")


def parse_timesheet_excel(path: str) -> tuple[str, list[dict[str, str] | None]]:
    """Read the company timesheet .xlsx and return (start_weekday, entries).

    entries follows the same shape as Row.to_entry(): None for an Off day,
    otherwise a dict with default clock times and Task/Detail mapped to
    Activity/Description.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    entries: list[dict[str, str] | None] = []
    start_weekday = ""
    row = TIMESHEET_DATA_START_ROW
    while True:
        no_val = ws.cell(row=row, column=COL_NO).value
        if no_val is None or str(no_val).strip() == "":
            break

        day_val = ws.cell(row=row, column=COL_DAY).value
        weekday = str(day_val).strip()[:3].title() if day_val else ""
        if not start_weekday and weekday in WEEKDAYS:
            start_weekday = weekday

        in_val = ws.cell(row=row, column=COL_IN).value
        out_val = ws.cell(row=row, column=COL_OUT).value
        if _is_zero_time(in_val) and _is_zero_time(out_val):
            entries.append(None)
        else:
            task_val = ws.cell(row=row, column=COL_TASK).value
            detail_val = ws.cell(row=row, column=COL_DETAIL).value
            entries.append({
                "editClockIn": join_time(DEFAULT_IN_TIME, DEFAULT_IN_AMPM),
                "editClockOut": join_time(DEFAULT_OUT_TIME, DEFAULT_OUT_AMPM),
                "editActivity": str(task_val).strip() if task_val else "",
                "editDescription": str(detail_val).strip() if detail_val else "",
            })
        row += 1

    return start_weekday or "Mon", entries


def split_time(value: str) -> tuple[str, str]:
    """'09:00 am' -> ('09:00', 'am'). Tolerant of missing/odd suffix."""
    parts = value.strip().rsplit(" ", 1)
    if len(parts) == 2 and parts[1].lower() in AMPM:
        return parts[0], parts[1].lower()
    return value.strip(), DEFAULT_IN_AMPM


def join_time(time_str: str, ampm: str) -> str:
    """('09:00', 'am') -> '09:00 am'."""
    return f"{time_str.strip()} {ampm}"


SCRIPT_TEMPLATE = """\
// =====================================================================
// LOGBOOK AUTO-FILL — generated by logbook_ui.py
// Paste this whole block into the browser console on the logbook page.
//   object -> fill row | null -> OFF day
// =====================================================================

const entries = {entries_json};

// =====================================================================
// One combined array, in document order = day order. Blue (empty) and orange
// (filled) day buttons both carry `.detailsbtn`, so this keeps index i aligned
// to day i regardless of color. Clicking a button edits that day whether it was
// blue or orange -> filling a filled day just overwrites it (auto-override).
const buttons = Array.from(document.querySelectorAll('.detailsbtn'));

buttons.forEach((b, i) => {{
    const entry = entries[i];

    if (entry && entry.skip) return; // leave this day untouched

    const editClockIn = document.getElementById('editClockIn');
    const editClockOut = document.getElementById('editClockOut');
    const editActivity = document.getElementById('editActivity');
    const editDescription = document.getElementById('editDescription');
    const submit = document.querySelector(
        '.button.button-primary[onclick="editSaveClick()"]'
    );
    const off = document.querySelector(
        '.button.button-primary[onclick="OffClick()"]'
    );
    b.click();

    if (entry === null || entry === undefined) {{
        off.click();
        submit.click();
        return;
    }}

    editClockIn.value = entry.editClockIn;
    editClockOut.value = entry.editClockOut;
    editActivity.value = entry.editActivity;
    editDescription.value = entry.editDescription;
    submit.click();
}});
"""


# ---- small custom widgets ---------------------------------------------------
class Field(tk.Frame):
    """Padded single-line entry with a 1px border that lights up on focus."""

    def __init__(
        self, parent: tk.Widget, app: "App", var: tk.StringVar, width: int
    ) -> None:
        p = app.palette
        super().__init__(
            parent, bg=p["input"], highlightthickness=1,
            highlightbackground=p["border"],
        )
        self.p = p
        self.entry = tk.Entry(
            self, textvariable=var, width=width, relief="flat", bd=0,
            bg=p["input"], fg=p["text"], insertbackground=p["text"],
            disabledbackground=p["input"], disabledforeground=p["dim"],
            highlightthickness=0, font=app.fonts["base"],
        )
        self.entry.pack(fill="x", expand=True, padx=6, pady=4)
        self.entry.bind("<FocusIn>", lambda e: self._focus(True))
        self.entry.bind("<FocusOut>", lambda e: self._focus(False))
        self.enabled = True

    def _focus(self, focused: bool) -> None:
        if self.enabled:
            self.config(
                highlightbackground=self.p["accent"] if focused else self.p["border"]
            )

    def set_enabled(self, enabled: bool, disabled_bg: str) -> None:
        self.enabled = enabled
        bg = self.p["input"] if enabled else disabled_bg
        self.config(bg=bg, highlightbackground=self.p["border"] if enabled else bg)
        self.entry.config(
            state="normal" if enabled else "disabled", disabledbackground=bg
        )


class Chip(tk.Label):
    """Clickable pill bound to a BooleanVar — replaces the plain checkbox."""

    def __init__(
        self, parent: tk.Widget, app: "App", text: str, var: tk.BooleanVar,
        on_color: str, command=None,
    ) -> None:
        super().__init__(
            parent, text=text, font=app.fonts["small_bold"], padx=10, pady=3,
            cursor="hand2",
        )
        self.app = app
        self.var = var
        self.on_color = on_color
        self.command = command
        self.enabled = True
        self.bind("<Button-1>", self._click)
        self.render()

    def _click(self, _event: tk.Event) -> None:
        if not self.enabled:
            return
        self.var.set(not self.var.get())
        self.render()
        if self.command:
            self.command()

    def set_enabled(self, enabled: bool) -> None:
        self.enabled = enabled
        self.config(cursor="hand2" if enabled else "")
        self.render()

    def render(self) -> None:
        p = self.app.palette
        if self.var.get():
            self.config(bg=self.on_color, fg=p["surface"])
        else:
            self.config(bg=p["chip"], fg=p["dim"] if self.enabled else p["faint"])


class AmPmToggle(tk.Label):
    """am/pm as a one-click toggle instead of a dropdown."""

    def __init__(self, parent: tk.Widget, app: "App", var: tk.StringVar) -> None:
        p = app.palette
        super().__init__(
            parent, textvariable=var, font=app.fonts["base"], width=3,
            padx=4, pady=4, bg=p["chip"], fg=p["text"], cursor="hand2",
        )
        self.p = p
        self.var = var
        self.enabled = True
        self.bind("<Button-1>", self._click)

    def _click(self, _event: tk.Event) -> None:
        if self.enabled:
            self.var.set("pm" if self.var.get() == "am" else "am")

    def set_enabled(self, enabled: bool, disabled_bg: str) -> None:
        self.enabled = enabled
        self.config(
            bg=self.p["chip"] if enabled else disabled_bg,
            fg=self.p["text"] if enabled else self.p["dim"],
            cursor="hand2" if enabled else "",
        )


def _focus_next(event: tk.Event) -> str:
    event.widget.tk_focusNext().focus_set()
    return "break"


def _focus_prev(event: tk.Event) -> str:
    event.widget.tk_focusPrev().focus_set()
    return "break"


class Row:
    """One logbook day = one row frame of widgets."""

    def __init__(
        self, parent: tk.Widget, index: int, weekday: str, app: "App"
    ) -> None:
        self.app = app
        self.index = index
        self.weekday = weekday
        self.locked = False
        p = app.palette

        self.off_var = tk.BooleanVar(value=False)
        self.skip_var = tk.BooleanVar(value=False)
        self.in_time = tk.StringVar(value=DEFAULT_IN_TIME)
        self.in_ampm = tk.StringVar(value=DEFAULT_IN_AMPM)
        self.out_time = tk.StringVar(value=DEFAULT_OUT_TIME)
        self.out_ampm = tk.StringVar(value=DEFAULT_OUT_AMPM)

        self.frame = tk.Frame(parent)
        configure_columns(self.frame)

        self.day_cell = tk.Frame(self.frame)
        self.day_label = tk.Label(
            self.day_cell, text=f"{index + 1}  {weekday}", anchor="w",
            font=app.fonts["bold"],
        )
        self.day_label.pack(anchor="w")
        self.day_note = tk.Label(self.day_cell, anchor="w", font=app.fonts["small"])

        self.off_chip = Chip(
            self.frame, app, "Off", self.off_var, p["dim"], command=self.refresh
        )
        self.skip_chip = Chip(
            self.frame, app, "Skip", self.skip_var, p["skip_fg"],
            command=self.refresh,
        )

        self.in_cell, self.in_field, self.in_toggle = self._time_cell(
            self.in_time, self.in_ampm
        )
        self.out_cell, self.out_field, self.out_toggle = self._time_cell(
            self.out_time, self.out_ampm
        )

        self.activity_e = self._text_cell()
        self.description_e = self._text_cell()

        cells = [
            (self.day_cell, "w"),
            (self.off_chip, "w"),
            (self.skip_chip, "w"),
            (self.in_cell, "w"),
            (self.out_cell, "w"),
            (self.activity_e, "ew"),
            (self.description_e, "ew"),
        ]
        for col, (w, sticky) in enumerate(cells):
            w.grid(row=0, column=col, padx=CELL_PADX, pady=6, sticky=sticky)

        self.refresh()

    def _time_cell(
        self, time_var: tk.StringVar, ampm_var: tk.StringVar
    ) -> tuple[tk.Frame, Field, AmPmToggle]:
        cell = tk.Frame(self.frame)
        field = Field(cell, self.app, time_var, width=6)
        field.pack(side="left")
        toggle = AmPmToggle(cell, self.app, ampm_var)
        toggle.pack(side="left", padx=(4, 0), fill="y")
        return cell, field, toggle

    def _text_cell(self) -> tk.Text:
        p = self.app.palette
        text = tk.Text(
            self.frame, width=1, height=2, wrap="word", relief="flat", bd=0,
            padx=6, pady=4, font=self.app.fonts["base"],
            bg=p["input"], fg=p["text"], insertbackground=p["text"],
            highlightthickness=1, highlightbackground=p["border"],
            highlightcolor=p["accent"],
        )
        text.bind("<MouseWheel>", self.app.forward_mousewheel)
        text.bind("<Button-4>", self.app.forward_mousewheel)
        text.bind("<Button-5>", self.app.forward_mousewheel)
        text.bind("<Tab>", _focus_next)
        text.bind("<Shift-Tab>", _focus_prev)
        text.bind("<KeyRelease>", lambda e: self.app.update_status())
        return text

    def pack(self) -> None:
        self.frame.pack(fill="x", pady=(0, 1))

    def destroy(self) -> None:
        self.frame.destroy()

    def get_activity(self) -> str:
        return self.activity_e.get("1.0", "end-1c")

    def set_activity(self, text: str) -> None:
        state = self.activity_e.cget("state")
        self.activity_e.config(state="normal")
        self.activity_e.delete("1.0", "end")
        self.activity_e.insert("1.0", text)
        self.activity_e.config(state=state)

    def get_description(self) -> str:
        return self.description_e.get("1.0", "end-1c")

    def set_description(self, text: str) -> None:
        state = self.description_e.cget("state")
        self.description_e.config(state="normal")
        self.description_e.delete("1.0", "end")
        self.description_e.insert("1.0", text)
        self.description_e.config(state=state)

    def refresh(self) -> None:
        """Repaint the row from its Off/Skip/Sunday state."""
        p = self.app.palette
        off, skip = self.off_var.get(), self.skip_var.get()
        sunday = self.weekday == "Sun"
        editable = not (off or skip)

        if sunday:
            bg = p["sun_bg"]
        elif skip:
            bg = p["skip_bg"]
        elif off:
            bg = p["off_bg"]
        else:
            bg = p["surface"] if self.index % 2 == 0 else p["surface_alt"]

        if sunday:
            fg, note = p["sun_fg"], "auto-off"
        elif skip:
            fg, note = p["skip_fg"], "untouched"
        elif off:
            fg, note = p["dim"], "day off"
        else:
            fg, note = p["text"], ""

        for w in (self.frame, self.day_cell, self.in_cell, self.out_cell):
            w.config(bg=bg)
        self.day_label.config(bg=bg, fg=fg)
        self.day_note.config(bg=bg, fg=fg, text=note)
        # only show the note line when there is one, so plain rows center the day
        if note:
            self.day_note.pack(anchor="w")
        else:
            self.day_note.pack_forget()

        self.off_chip.set_enabled(not self.locked and not skip)
        self.skip_chip.set_enabled(not self.locked)

        for field in (self.in_field, self.out_field):
            field.set_enabled(editable, bg)
        for toggle in (self.in_toggle, self.out_toggle):
            toggle.set_enabled(editable, bg)
        for text in (self.activity_e, self.description_e):
            text.config(
                state="normal" if editable else "disabled",
                bg=p["input"] if editable else bg,
                fg=p["text"] if editable else p["dim"],
                highlightbackground=p["border"] if editable else bg,
            )
        self.app.update_status()

    def lock_sunday(self) -> None:
        """Force this row OFF and lock its flags (Sundays have no site button)."""
        self.off_var.set(True)
        self.skip_var.set(False)
        self.locked = True
        self.refresh()

    def to_entry(self) -> dict[str, str] | None:
        if self.skip_var.get():
            return {"skip": True}
        if self.off_var.get():
            return None
        return {
            "editClockIn": join_time(self.in_time.get(), self.in_ampm.get()),
            "editClockOut": join_time(self.out_time.get(), self.out_ampm.get()),
            "editActivity": self.get_activity(),
            "editDescription": self.get_description(),
        }

    def load(self, entry: dict[str, str] | None) -> None:
        if entry is None:
            self.off_var.set(True)
            self.skip_var.set(False)
            self.refresh()
            return
        if entry.get("skip"):
            self.off_var.set(False)
            self.skip_var.set(True)
            self.refresh()
            return
        self.off_var.set(False)
        self.skip_var.set(False)
        self.refresh()
        it, ia = split_time(entry.get("editClockIn", DEFAULT_IN_TIME))
        ot, oa = split_time(entry.get("editClockOut", DEFAULT_OUT_TIME))
        self.in_time.set(it)
        self.in_ampm.set(ia)
        self.out_time.set(ot)
        self.out_ampm.set(oa)
        self.set_activity(entry.get("editActivity", ""))
        self.set_description(entry.get("editDescription", ""))


class App:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        root.title("Logbook Entry Editor")
        self.palette = DARK if detect_dark_mode() else LIGHT
        self.rows: list[Row] = []
        self._toast_job: str | None = None

        self._setup_style()
        self._build_header()
        self._build_bulk_panel()
        self._build_footer()  # packed bottom-first so it never gets squeezed
        self._build_table()

        self.rebuild_rows()

    # ---- theme -----------------------------------------------------------
    def _setup_style(self) -> None:
        p = self.palette
        family = tkfont.nametofont("TkDefaultFont").actual("family")
        size = 13 if sys.platform == "darwin" else 10
        self.fonts = {
            "base": (family, size),
            "bold": (family, size, "bold"),
            "small": (family, size - 2),
            "small_bold": (family, size - 2, "bold"),
            "title": (family, size + 3, "bold"),
        }
        self.root.configure(bg=p["bg"])

        style = ttk.Style(self.root)
        style.theme_use("clam")
        style.configure(
            ".", background=p["surface"], foreground=p["text"],
            font=self.fonts["base"], bordercolor=p["border"],
            focuscolor=p["accent"],
        )

        flat = {"bordercolor": p["border"], "relief": "flat"}
        style.configure(
            "TButton", background=p["surface"], foreground=p["text"],
            lightcolor=p["surface"], darkcolor=p["surface"], padding=(12, 5),
            **flat,
        )
        hover = [("pressed", p["border"]), ("active", p["chip"])]
        style.map("TButton", background=hover, lightcolor=hover, darkcolor=hover)

        style.configure(
            "Accent.TButton", background=p["accent"], foreground=p["accent_fg"],
            bordercolor=p["accent"], lightcolor=p["accent"], darkcolor=p["accent"],
            font=self.fonts["bold"], padding=(16, 6),
        )
        accent_hover = [("pressed", p["accent_hover"]), ("active", p["accent_hover"])]
        style.map(
            "Accent.TButton", background=accent_hover, lightcolor=accent_hover,
            darkcolor=accent_hover, bordercolor=accent_hover,
        )

        field = {
            "fieldbackground": p["input"], "background": p["chip"],
            "foreground": p["text"], "arrowcolor": p["dim"],
            "lightcolor": p["input"], "darkcolor": p["input"],
            "insertcolor": p["text"], "padding": 4, **flat,
        }
        style.configure("TSpinbox", **field)
        style.configure("TCombobox", **field)
        style.map(
            "TCombobox",
            fieldbackground=[("readonly", p["input"])],
            foreground=[("readonly", p["text"])],
            selectbackground=[("readonly", p["input"])],
            selectforeground=[("readonly", p["text"])],
        )
        focus_border = [("focus", p["accent"])]
        style.map("TSpinbox", bordercolor=focus_border)
        self.root.option_add("*TCombobox*Listbox.background", p["surface"])
        self.root.option_add("*TCombobox*Listbox.foreground", p["text"])
        self.root.option_add("*TCombobox*Listbox.selectBackground", p["accent"])
        self.root.option_add("*TCombobox*Listbox.selectForeground", p["accent_fg"])

        style.configure(
            "Vertical.TScrollbar", background=p["chip"], troughcolor=p["surface"],
            bordercolor=p["surface"], arrowcolor=p["dim"],
            lightcolor=p["chip"], darkcolor=p["chip"],
        )

    def _label(
        self, parent: tk.Widget, text: str, bg: str, font: str = "base",
        fg: str | None = None,
    ) -> tk.Label:
        return tk.Label(
            parent, text=text, bg=bg, fg=fg or self.palette["dim"],
            font=self.fonts[font],
        )

    # ---- UI construction -------------------------------------------------
    def _build_header(self) -> None:
        p = self.palette
        bar = tk.Frame(self.root, bg=p["surface"])
        bar.pack(fill="x")
        tk.Frame(self.root, bg=p["border"], height=1).pack(fill="x")
        inner = tk.Frame(bar, bg=p["surface"])
        inner.pack(fill="x", padx=16, pady=12)

        # triangle logo (matches web/index.html); Tk has no gradients, so solid accent
        logo = tk.Canvas(
            inner, width=36, height=36, bg=p["surface"], highlightthickness=0
        )
        logo.create_polygon(
            18, 5, 32, 30, 4, 30, fill=p["accent"], outline=p["accent"],
            width=4, joinstyle="round",
        )
        logo.create_polygon(
            18, 15, 24, 26, 12, 26, fill="", outline=p["surface"], width=2,
            joinstyle="round",
        )
        logo.pack(side="left")
        titles = tk.Frame(inner, bg=p["surface"])
        titles.pack(side="left", padx=(10, 0))
        self._label(
            titles, "Logbook Entry Editor", p["surface"], "title", p["text"]
        ).pack(anchor="w")
        self._label(
            titles,
            "Build the month here, generate the auto-fill script, "
            "paste it into the logbook console.",
            p["surface"], "small",
        ).pack(anchor="w")

        # right side, packed right-to-left
        ttk.Button(inner, text="Rebuild rows", command=self.rebuild_rows).pack(
            side="right", padx=(10, 0)
        )
        self.start_var = tk.StringVar(value="Mon")
        ttk.Combobox(
            inner, textvariable=self.start_var, values=WEEKDAYS, width=5,
            state="readonly",
        ).pack(side="right")
        self._label(inner, "Start weekday", p["surface"]).pack(
            side="right", padx=(14, 6)
        )
        self.days_var = tk.IntVar(value=28)
        ttk.Spinbox(
            inner, from_=1, to=31, textvariable=self.days_var, width=4
        ).pack(side="right")
        self._label(inner, "Days", p["surface"]).pack(side="right", padx=(0, 6))

    def _build_bulk_panel(self) -> None:
        p = self.palette
        bg = p["surface"]
        card = tk.Frame(
            self.root, bg=bg, highlightthickness=1, highlightbackground=p["border"]
        )
        card.pack(fill="x", padx=16, pady=(12, 0))

        head = tk.Frame(card, bg=bg, cursor="hand2")
        head.pack(fill="x", padx=14, pady=10)
        self.bulk_arrow = self._label(head, "▾", bg, "bold", p["text"])
        self.bulk_arrow.pack(side="left")
        title = self._label(head, "Bulk tools", bg, "bold", p["text"])
        title.pack(side="left", padx=(6, 0))
        hint = self._label(
            head, "Applies to every row that isn't Off or Skip", bg, "small"
        )
        hint.pack(side="right")
        for w in (head, self.bulk_arrow, title, hint):
            w.bind("<Button-1>", lambda e: self.toggle_bulk())

        body = tk.Frame(card, bg=bg)
        body.grid_columnconfigure(1, weight=1)
        tk.Frame(body, bg=p["border"], height=1).grid(
            row=0, column=0, columnspan=3, sticky="ew", pady=(0, 8)
        )
        self.bulk_body = body

        def section(row: int, name: str) -> tk.Frame:
            self._label(body, name.upper(), bg, "small_bold").grid(
                row=row, column=0, sticky="w", padx=(0, 16), pady=5
            )
            line = tk.Frame(body, bg=bg)
            line.grid(row=row, column=1, sticky="ew", pady=5)
            return line

        def apply_button(row: int, text: str, command) -> None:
            ttk.Button(body, text=text, command=command).grid(
                row=row, column=2, sticky="ew", padx=(16, 0), pady=5
            )

        # ---- bulk clock-time setter ----
        line = section(1, "Time")
        self.bulk_in_time = tk.StringVar(value=DEFAULT_IN_TIME)
        self.bulk_in_ampm = tk.StringVar(value=DEFAULT_IN_AMPM)
        self.bulk_out_time = tk.StringVar(value=DEFAULT_OUT_TIME)
        self.bulk_out_ampm = tk.StringVar(value=DEFAULT_OUT_AMPM)
        self._label(line, "In", bg).pack(side="left", padx=(0, 6))
        Field(line, self, self.bulk_in_time, width=6).pack(side="left")
        AmPmToggle(line, self, self.bulk_in_ampm).pack(
            side="left", padx=(4, 0), fill="y"
        )
        self._label(line, "Out", bg).pack(side="left", padx=(18, 6))
        Field(line, self, self.bulk_out_time, width=6).pack(side="left")
        AmPmToggle(line, self, self.bulk_out_ampm).pack(
            side="left", padx=(4, 0), fill="y"
        )
        apply_button(1, "Apply times", self.apply_clock_all)

        # ---- bulk activity / description setter ----
        line = section(2, "Text")
        self.bulk_activity = tk.StringVar(value="")
        self.bulk_description = tk.StringVar(value="")
        self._label(line, "Activity", bg).pack(side="left", padx=(0, 6))
        Field(line, self, self.bulk_activity, width=20).pack(
            side="left", fill="x", expand=True
        )
        self._label(line, "Description", bg).pack(side="left", padx=(18, 6))
        Field(line, self, self.bulk_description, width=24).pack(
            side="left", fill="x", expand=True
        )
        apply_button(2, "Apply text", self.apply_text_all)

        # ---- subheader prepend ----
        line = section(3, "Subheader")
        self.subheader_var = tk.StringVar(value="")
        self.subheader_from_var = tk.IntVar(value=1)
        self.subheader_to_var = tk.IntVar(value=1)
        Field(line, self, self.subheader_var, width=20).pack(
            side="left", fill="x", expand=True
        )
        self._label(line, "From day", bg).pack(side="left", padx=(18, 6))
        ttk.Spinbox(
            line, from_=1, to=31, textvariable=self.subheader_from_var, width=4
        ).pack(side="left")
        self._label(line, "To day", bg).pack(side="left", padx=(12, 6))
        ttk.Spinbox(
            line, from_=1, to=31, textvariable=self.subheader_to_var, width=4
        ).pack(side="left")
        apply_button(3, "Apply subheader", self.apply_subheader)
        self._label(
            body,
            "Prepends to Activity, e.g. (XXXXX Project) Ticket AX-122. "
            "Re-applying stacks it again.",
            bg, "small",
        ).grid(row=4, column=1, sticky="w")

        self.bulk_open = False
        self.toggle_bulk()

    def toggle_bulk(self) -> None:
        self.bulk_open = not self.bulk_open
        if self.bulk_open:
            self.bulk_body.pack(fill="x", padx=14, pady=(0, 12))
        else:
            self.bulk_body.pack_forget()
        self.bulk_arrow.config(text="▾" if self.bulk_open else "▸")

    def _build_table(self) -> None:
        p = self.palette
        card = tk.Frame(
            self.root, bg=p["surface"], highlightthickness=1,
            highlightbackground=p["border"],
        )
        card.pack(fill="both", expand=True, padx=16, pady=(12, 0))
        card.grid_rowconfigure(2, weight=1)
        card.grid_columnconfigure(0, weight=1)

        # header lives outside the canvas so it stays put while rows scroll
        header = tk.Frame(card, bg=p["surface_alt"])
        configure_columns(header)
        header.grid(row=0, column=0, sticky="ew")
        for col, (title, _, _) in enumerate(COLUMNS):
            self._label(header, title.upper(), p["surface_alt"], "small_bold").grid(
                row=0, column=col, padx=CELL_PADX, pady=8, sticky="w"
            )
        tk.Frame(card, bg=p["surface_alt"]).grid(row=0, column=1, sticky="nsew")
        tk.Frame(card, bg=p["border"], height=1).grid(
            row=1, column=0, columnspan=2, sticky="ew"
        )

        canvas = tk.Canvas(card, highlightthickness=0, bd=0, bg=p["surface"])
        vbar = ttk.Scrollbar(card, orient="vertical", command=canvas.yview)
        canvas.grid(row=2, column=0, sticky="nsew")
        vbar.grid(row=2, column=1, sticky="ns")

        # border-colored backdrop + 1px gap under each row = row separators
        self.grid_frame = tk.Frame(canvas, bg=p["border"])
        window = canvas.create_window((0, 0), window=self.grid_frame, anchor="nw")
        self.grid_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all")),
        )
        canvas.bind("<Configure>", lambda e: canvas.itemconfigure(window, width=e.width))
        canvas.configure(yscrollcommand=vbar.set)

        self.canvas = canvas
        canvas.bind_all("<MouseWheel>", self.forward_mousewheel)
        # X11 (Linux) sends Button-4/5 wheel events instead of <MouseWheel>
        canvas.bind_all("<Button-4>", self.forward_mousewheel)
        canvas.bind_all("<Button-5>", self.forward_mousewheel)

    def _build_footer(self) -> None:
        p = self.palette
        status = tk.Frame(self.root, bg=p["bg"])
        status.pack(side="bottom", fill="x", padx=18, pady=(0, 10))
        self.status_label = self._label(status, "", p["bg"], "small")
        self.status_label.pack(side="left")
        self.toast_label = self._label(status, "", p["bg"], "small_bold", p["accent"])
        self.toast_label.pack(side="right")

        bar = tk.Frame(self.root, bg=p["bg"])
        bar.pack(side="bottom", fill="x", padx=16, pady=(12, 8))
        for text, command in (
            ("Load JSON…", self.load_json),
            ("Save JSON…", self.save_json),
            ("Import Excel…", self.import_excel),
            ("Reset All", self.reset_all),
        ):
            ttk.Button(bar, text=text, command=command).pack(side="left", padx=(0, 6))
        ttk.Button(
            bar, text="Generate script.js + copy", style="Accent.TButton",
            command=self.generate,
        ).pack(side="right")

    # ---- feedback ----------------------------------------------------------
    def update_status(self) -> None:
        rows = self.rows
        working = [r for r in rows if not r.off_var.get() and not r.skip_var.get()]
        filled = sum(1 for r in working if r.get_activity().strip())
        off = sum(1 for r in rows if r.off_var.get())
        skip = sum(1 for r in rows if r.skip_var.get())
        sundays = sum(1 for r in rows if r.weekday == "Sun")
        self.status_label.config(
            text=f"{len(rows)} days   ·   {off} off ({sundays} Sun)   ·   "
            f"{skip} skip   ·   {filled}/{len(working)} filled"
        )

    def toast(self, message: str, warn: bool = False) -> None:
        """Show a short message in the status bar that clears itself."""
        p = self.palette
        self.toast_label.config(
            text=message, fg=p["sun_fg"] if warn else p["accent"]
        )
        if self._toast_job:
            self.root.after_cancel(self._toast_job)
        self._toast_job = self.root.after(
            4000, lambda: self.toast_label.config(text="")
        )

    # ---- row management --------------------------------------------------
    def rebuild_rows(self) -> None:
        # preserve existing values where possible
        old = [r.to_entry() for r in self.rows]
        for r in self.rows:
            r.destroy()
        self.rows.clear()

        n = max(1, self.days_var.get())
        start = WEEKDAYS.index(self.start_var.get())
        for i in range(n):
            row = Row(self.grid_frame, i, WEEKDAYS[(start + i) % 7], self)
            row.pack()
            if i < len(old):
                row.load(old[i])
            self.rows.append(row)
        self._lock_sundays()
        self.update_status()

    def _lock_sundays(self) -> None:
        for row in self.rows:
            if row.weekday == "Sun":
                row.lock_sunday()

    def forward_mousewheel(self, event: tk.Event) -> str:
        """Scroll the row table on any mouse wheel event, even over a Text box.

        macOS reports small per-notch deltas (no /120 scaling like Windows),
        so the two platforms need different unit conversions.
        """
        if event.num == 4:
            self.canvas.yview_scroll(-3, "units")
        elif event.num == 5:
            self.canvas.yview_scroll(3, "units")
        elif sys.platform == "darwin":
            self.canvas.yview_scroll(int(-event.delta), "units")
        else:
            self.canvas.yview_scroll(int(-event.delta / 120), "units")
        return "break"

    def _editable_rows(self) -> list[Row]:
        return [r for r in self.rows if not r.off_var.get() and not r.skip_var.get()]

    def apply_clock_all(self) -> None:
        """Set clock in/out on every non-off row to the bulk values."""
        rows = self._editable_rows()
        for row in rows:
            row.in_time.set(self.bulk_in_time.get())
            row.in_ampm.set(self.bulk_in_ampm.get())
            row.out_time.set(self.bulk_out_time.get())
            row.out_ampm.set(self.bulk_out_ampm.get())
        self.toast(f"Times applied to {len(rows)} rows")

    def reset_all(self) -> None:
        """Restore every row to defaults (times + activity/description, clear flags)."""
        if not messagebox.askyesno(
            "Reset All",
            "Reset every row's clock times and activity/description to defaults "
            "and clear all Off/Skip flags?\n\nThis cannot be undone.",
        ):
            return
        for row in self.rows:
            row.off_var.set(False)
            row.skip_var.set(False)
            row.in_time.set(DEFAULT_IN_TIME)
            row.in_ampm.set(DEFAULT_IN_AMPM)
            row.out_time.set(DEFAULT_OUT_TIME)
            row.out_ampm.set(DEFAULT_OUT_AMPM)
            row.refresh()  # re-enable fields before clearing the text
            row.set_activity("")
            row.set_description("")
        self._lock_sundays()
        self.update_status()
        self.toast("All rows reset")

    def apply_text_all(self) -> None:
        """Set activity/description on every non-off row to the bulk values."""
        rows = self._editable_rows()
        for row in rows:
            row.set_activity(self.bulk_activity.get())
            row.set_description(self.bulk_description.get())
        self.update_status()
        self.toast(f"Text applied to {len(rows)} rows")

    def apply_subheader(self) -> None:
        """Prepend '<subheader> ' to Activity for rows in the day range."""
        subheader = self.subheader_var.get().strip()
        if not subheader:
            self.toast("Enter a subheader first", warn=True)
            return
        start = min(self.subheader_from_var.get(), self.subheader_to_var.get())
        end = max(self.subheader_from_var.get(), self.subheader_to_var.get())
        count = 0
        for i, row in enumerate(self.rows, start=1):
            if i < start or i > end:
                continue
            if row.off_var.get() or row.skip_var.get():
                continue
            row.set_activity(f"{subheader} {row.get_activity()}".rstrip())
            count += 1
        self.update_status()
        self.toast(f"Subheader added to {count} rows (days {start}–{end})")

    # ---- import / export -------------------------------------------------
    def collect(self) -> list[dict[str, str] | None]:
        return [r.to_entry() for r in self.rows]

    def load_json(self) -> None:
        path = filedialog.askopenfilename(
            filetypes=[("JSON", "*.json"), ("All", "*.*")]
        )
        if not path:
            return
        data = json.loads(Path(path).read_text(encoding="utf-8"))
        self.days_var.set(len(data))
        self.rebuild_rows()
        for row, entry in zip(self.rows, data):
            row.load(entry)
        self._lock_sundays()
        self.toast(f"Loaded {len(data)} rows from {Path(path).name}")

    def import_excel(self) -> None:
        path = filedialog.askopenfilename(
            filetypes=[("Excel", "*.xlsx"), ("All", "*.*")]
        )
        if not path:
            return
        try:
            start_weekday, entries = parse_timesheet_excel(path)
        except Exception as e:
            messagebox.showerror("Import failed", str(e))
            return
        if not entries:
            self.toast(
                f"No data rows found starting at row {TIMESHEET_DATA_START_ROW}",
                warn=True,
            )
            return

        self.days_var.set(len(entries))
        self.start_var.set(start_weekday)
        self.rebuild_rows()
        for row, entry in zip(self.rows, entries):
            row.load(entry)
        self._lock_sundays()
        self.toast(f"Imported {len(entries)} rows from {Path(path).name}")

    def save_json(self) -> None:
        path = filedialog.asksaveasfilename(
            defaultextension=".json",
            filetypes=[("JSON", "*.json")],
        )
        if not path:
            return
        Path(path).write_text(
            json.dumps(self.collect(), indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
        self.toast(f"Saved {Path(path).name}")

    def generate(self) -> None:
        # Sundays have no button on the site (auto-OFF), so drop them here to
        # keep entries[] aligned with the buttons the page actually renders.
        rows = [r for r in self.rows if r.weekday != "Sun"]
        skipped_sundays = len(self.rows) - len(rows)
        entries = [r.to_entry() for r in rows]
        entries_json = json.dumps(entries, indent=4, ensure_ascii=False)
        script = SCRIPT_TEMPLATE.format(entries_json=entries_json)

        out = Path(__file__).with_name("script.generated.js")
        out.write_text(script, encoding="utf-8")

        self.root.clipboard_clear()
        self.root.clipboard_append(script)
        self.toast(
            f"Copied to clipboard + saved {out.name}: {len(entries)} rows "
            f"({skipped_sundays} Sundays excluded)"
        )


def main() -> None:
    root = tk.Tk()
    App(root)
    root.geometry("1100x760")
    root.minsize(900, 560)
    root.mainloop()


if __name__ == "__main__":
    main()
